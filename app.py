import io
import re
from datetime import datetime
from pathlib import Path

import pandas as pd
import geopandas as gpd
import streamlit as st
from openpyxl import load_workbook

try:
    from zoneinfo import ZoneInfo
except ImportError:  # pragma: no cover
    ZoneInfo = None

# =========================
# PAGE CONFIG
# =========================
st.set_page_config(page_title="Ticket Automation", page_icon="🎫", layout="centered")

# =========================
# FILE PATHS (relative to this file's location)
# =========================
BASE_DIR = Path(__file__).resolve().parent

# NOTE: The local "Input CSV" folder is no longer used.
# CSV files are now uploaded by the user via the Streamlit web interface.

# Shapefile base name and the companion files that must exist next to the .shp
SHAPEFILE_NAME = "gadm41_IDN_3"
SHAPEFILE_EXTS = [".shp", ".shx", ".dbf", ".prj"]

# Indonesian month names
bulan_id = {
    1: "Januari",
    2: "Februari",
    3: "Maret",
    4: "April",
    5: "Mei",
    6: "Juni",
    7: "Juli",
    8: "Agustus",
    9: "September",
    10: "Oktober",
    11: "November",
    12: "Desember"
}

WAKTU_COL = 6  # Column F

# How the column should look in the final Excel file — minutes only, no seconds
OUTPUT_NUMBER_FORMAT = "dd/mm/yyyy hh:mm"

# Formats we'll try, in order, to parse the raw source text
CANDIDATE_FORMATS = [
    "%d/%m/%Y %H:%M:%S",
    "%d/%m/%Y %H:%M",
    "%Y-%m-%d %H:%M:%S",
    "%Y-%m-%d %H:%M:%S.%f",
    "%Y-%m-%d %H:%M",
    "%d-%m-%Y %H:%M:%S",
    "%m/%d/%Y %H:%M:%S",
]


class InvalidCSVError(Exception):
    """Raised when an uploaded CSV file is invalid or malformed."""


class MissingShapefileError(Exception):
    """Raised when the shapefile (or one of its companion files) is missing."""


def get_today():
    """
    Returns the current datetime. Uses Asia/Jakarta time zone when available
    (Streamlit Cloud servers run in UTC), otherwise falls back to server time.
    """
    if ZoneInfo is not None:
        try:
            return datetime.now(ZoneInfo("Asia/Jakarta"))
        except Exception:
            pass
    return datetime.now()


def build_output_filename():
    """Tiket Indihome DD MMMM YYYY butuh feedback.xlsx"""
    today = get_today()
    tanggal_indonesia = f"{today.day:02d} {bulan_id[today.month]} {today.year}"
    return f"Tiket Indihome {tanggal_indonesia} butuh feedback.xlsx"


def manual_parse_csv_text(raw_text):
    """
    Pure character-by-character CSV parser (no csv module, no pandas).
    Respects quoted fields so commas inside quotes don't break column
    alignment. Performs zero interpretation of values.
    """
    rows = []
    field = []
    row = []
    in_quotes = False
    i = 0
    n = len(raw_text)

    while i < n:
        ch = raw_text[i]

        if in_quotes:
            if ch == '"':
                if i + 1 < n and raw_text[i + 1] == '"':
                    field.append('"')
                    i += 2
                    continue
                in_quotes = False
                i += 1
                continue
            field.append(ch)
            i += 1
            continue

        if ch == '"':
            in_quotes = True
            i += 1
            continue
        elif ch == ",":
            row.append("".join(field))
            field = []
            i += 1
            continue
        elif ch == "\r":
            i += 1
            continue
        elif ch == "\n":
            row.append("".join(field))
            field = []
            rows.append(row)
            row = []
            i += 1
            continue
        else:
            field.append(ch)
            i += 1
            continue

    if field or row:
        row.append("".join(field))
        rows.append(row)

    return rows


def parse_waktu_value(raw_value):
    """
    Converts raw column F text into a real Python datetime object.
    Tries several known formats explicitly. Falls back to pandas'
    flexible parser (dayfirst=True) if none match. Returns the
    original raw string only if every attempt fails.
    """
    cleaned = re.sub(r"\s+", " ", str(raw_value).strip())

    if not cleaned:
        return raw_value

    for fmt in CANDIDATE_FORMATS:
        try:
            return datetime.strptime(cleaned, fmt)
        except ValueError:
            continue

    # Fallback: flexible parsing, day-first (26/09/2026 style)
    try:
        parsed = pd.to_datetime(cleaned, dayfirst=True, errors="raise")
        return parsed.to_pydatetime()
    except Exception:
        return raw_value


# =========================
# CHECK REQUIRED SHAPEFILE
# =========================
def check_shapefile():
    """
    Locates the shapefile anywhere inside the repository (case-insensitive)
    and verifies that all companion files exist next to it.
    Returns the Path to the .shp file.
    """
    search_roots = [BASE_DIR]
    if Path.cwd().resolve() != BASE_DIR:
        search_roots.append(Path.cwd().resolve())

    shp_path = None
    for root in search_roots:
        for p in root.rglob("*"):
            if ".git" in p.parts:
                continue
            if p.is_file() and p.name.lower() == f"{SHAPEFILE_NAME}.shp".lower():
                shp_path = p
                break
        if shp_path:
            break

    if shp_path is None:
        # Build a small diagnostic so the real repo layout is visible
        listing = []
        for root in search_roots:
            for p in sorted(root.rglob("*")):
                if ".git" in p.parts or "__pycache__" in p.parts:
                    continue
                listing.append(str(p.relative_to(root)))
                if len(listing) >= 60:
                    break
        raise MissingShapefileError(
            f"Could not find '{SHAPEFILE_NAME}.shp' anywhere in the repository.\n"
            f"App location: {BASE_DIR}\n"
            f"Files visible to the app:\n" + "\n".join(listing)
        )

    # Verify companion files in the same folder (case-insensitive)
    names_in_dir = {p.name.lower() for p in shp_path.parent.iterdir()}
    missing = [
        f"{SHAPEFILE_NAME}{ext}"
        for ext in SHAPEFILE_EXTS
        if f"{SHAPEFILE_NAME}{ext}".lower() not in names_in_dir
    ]
    if missing:
        raise MissingShapefileError(
            f"Found {shp_path.name} in '{shp_path.parent}', but these companion "
            f"files are missing: {', '.join(missing)}"
        )

    return shp_path


@st.cache_resource(show_spinner=False)
def load_admin_shapefile(shp_path_str):
    """Reads the shapefile once and caches it across reruns."""
    return gpd.read_file(shp_path_str)


# =========================
# MAIN PROCESSING FUNCTION
# =========================
def process_uploaded_files(uploaded_files):
    """
    Runs the full pipeline on the uploaded CSV files and returns
    (excel_bytes, stats_dict). Nothing is written to disk.
    """

    # =========================
    # CHECK REQUIRED FILES
    # =========================
    shp_path = check_shapefile()

    # =========================
    # FIND CSV FILES (uploaded via Streamlit)
    # =========================
    csv_files = sorted(
        (f for f in uploaded_files if f.name.lower().endswith(".csv")),
        key=lambda f: f.name.lower(),
    )

    if not csv_files:
        raise ValueError("No CSV files uploaded.")

    # =========================
    # SAVE ORIGINAL WAKTU COLUMN (F) AS PARSED DATETIME
    # =========================

    waktu_data = []
    waktu_col_name = None
    all_dfs = []

    for csv_file in csv_files:

        print(f"Reading: {csv_file.name}")

        file_bytes = csv_file.getvalue()

        try:
            raw_text = file_bytes.decode("utf-8-sig")
        except UnicodeDecodeError as e:
            raise InvalidCSVError(
                f"File '{csv_file.name}' is not a valid UTF-8 CSV file: {e}"
            )

        parsed_rows = manual_parse_csv_text(raw_text)

        if not parsed_rows:
            raise InvalidCSVError(f"File '{csv_file.name}' appears to be empty.")

        header = parsed_rows[0]

        if len(header) < WAKTU_COL:
            raise InvalidCSVError(
                f"File '{csv_file.name}' has fewer than {WAKTU_COL} columns."
            )

        this_waktu_col_name = header[WAKTU_COL - 1].strip()

        if waktu_col_name is None:
            waktu_col_name = this_waktu_col_name

        for data_row in parsed_rows[1:]:
            raw_value = data_row[WAKTU_COL - 1] if len(data_row) > WAKTU_COL - 1 else ""
            waktu_data.append(parse_waktu_value(raw_value))

        # Read the SAME file with pandas, EXCLUDING column F entirely
        try:
            temp_df = pd.read_csv(
                io.BytesIO(file_bytes),
                dtype=str,
                keep_default_na=False,
                low_memory=False,
                encoding="utf-8-sig",
                usecols=lambda col: col.strip() != this_waktu_col_name,
            )
        except Exception as e:
            raise InvalidCSVError(
                f"File '{csv_file.name}' could not be read as a CSV: {e}"
            )

        temp_df.columns = temp_df.columns.astype(str).str.strip()

        all_dfs.append(temp_df)

    print("DEBUG - first 3 waktu_data values:", waktu_data[:3])
    print("DEBUG - types:", [type(v) for v in waktu_data[:3]])

    # =========================
    # MERGE CSV FILES (WITHOUT COLUMN F)
    # =========================
    df = pd.concat(all_dfs, ignore_index=True)
    df.columns = df.columns.astype(str).str.strip()

    # =========================
    # ENSURE REQUIRED COLUMNS
    # =========================
    if "Grid GeoJSON" not in df.columns:
        raise InvalidCSVError("Column 'Grid GeoJSON' not found.")

    if "Mapping RCA" not in df.columns:
        print("'Mapping RCA' not found. Creating column...")
        df["Mapping RCA"] = ""

    # =========================
    # KEEP ONLY COLUMNS FROM A UNTIL MAPPING RCA
    # =========================
    mapping_idx = df.columns.get_loc("Mapping RCA")
    keep_cols = list(df.columns[:mapping_idx + 1])

    if "Grid GeoJSON" not in keep_cols:
        keep_cols.append("Grid GeoJSON")

    df = df[keep_cols]

    # =========================
    # EXTRACT LONGITUDE & LATITUDE
    # =========================
    df["Longitude"] = df["Grid GeoJSON"].astype(str).str.extract(
        r"\[\[\[\[\s*([-\d\.]+)"
    )
    df["Latitude"] = df["Grid GeoJSON"].astype(str).str.extract(
        r"\[\[\[\[\s*[-\d\.]+\s*,\s*([-\d\.]+)"
    )

    df["Longitude"] = pd.to_numeric(df["Longitude"], errors="coerce")
    df["Latitude"] = pd.to_numeric(df["Latitude"], errors="coerce")

    # =========================
    # CREATE GEODATAFRAME
    # =========================
    gdf_points = gpd.GeoDataFrame(
        df.copy(),
        geometry=gpd.points_from_xy(df["Longitude"], df["Latitude"]),
        crs="EPSG:4326"
    )

    # =========================
    # READ SHAPEFILE
    # =========================
    gdf_admin = load_admin_shapefile(str(shp_path))

    if gdf_admin.crs != gdf_points.crs:
        gdf_admin = gdf_admin.to_crs(gdf_points.crs)

    # =========================
    # SPATIAL JOIN
    # =========================
    result = gpd.sjoin(
        gdf_points,
        gdf_admin,
        how="left",
        predicate="within"
    )

    # =========================
    # FALLBACK FOR BOUNDARY POINTS
    # =========================
    missing = result["NAME_1"].isna()

    if missing.any():
        fallback = gpd.sjoin_nearest(
            gdf_points[missing],
            gdf_admin,
            how="left"
        )
        result.loc[missing, ["NAME_1", "NAME_2"]] = fallback[["NAME_1", "NAME_2"]].values

    # =========================
    # ADD PROVINSI & KABUPATEN
    # =========================
    df["Provinsi"] = result["NAME_1"]
    df["Kabupaten/Kota"] = result["NAME_2"]

    # =========================
    # INSERT AFTER MAPPING RCA
    # =========================
    mapping_idx = df.columns.get_loc("Mapping RCA")

    provinsi = df.pop("Provinsi")
    kabupaten = df.pop("Kabupaten/Kota")

    df.insert(mapping_idx + 1, "Provinsi", provinsi)
    df.insert(mapping_idx + 2, "Kabupaten/Kota", kabupaten)

    # =========================
    # RE-INSERT WAKTU AS BLANK PLACEHOLDER
    # =========================
    row_count_warning = None

    if len(waktu_data) != len(df):
        row_count_warning = (
            "WARNING: number of column F values "
            f"({len(waktu_data)}) does not match number of "
            f"output rows ({len(df)}). Check source CSVs for "
            "blank lines or ragged rows."
        )
        print(row_count_warning)

    insert_pos = min(WAKTU_COL - 1, len(df.columns))
    df.insert(insert_pos, waktu_col_name, "")

    # =========================
    # SAVE RESULT (in memory, no file written to disk)
    # =========================
    temp_buffer = io.BytesIO()
    df.to_excel(
        temp_buffer,
        index=False
    )
    temp_buffer.seek(0)

    # =========================
    # OPEN OUTPUT FILE
    # =========================
    wb_output = load_workbook(temp_buffer)
    ws_output = wb_output.active

    # =========================
    # CLEAR COLUMN F
    # =========================
    for row in range(2, ws_output.max_row + 1):
        ws_output.cell(row=row, column=WAKTU_COL).value = None

    # =========================
    # RESTORE COLUMN F AS REAL DATETIME
    # =========================
    converted_count = 0
    fallback_count = 0

    for row, value in enumerate(waktu_data, start=2):

        if row > ws_output.max_row:
            break

        tgt = ws_output.cell(row=row, column=WAKTU_COL)
        tgt.value = value

        if isinstance(value, datetime):
            tgt.number_format = OUTPUT_NUMBER_FORMAT
            converted_count += 1
        else:
            tgt.number_format = "@"
            fallback_count += 1

    # =========================
    # SAVE FINAL FILE (in memory)
    # =========================
    final_buffer = io.BytesIO()
    wb_output.save(final_buffer)
    final_buffer.seek(0)

    print("Waktu column converted to real datetime values.")
    print("Completed!")
    print(f"Rows processed : {len(df)}")
    print(f"CSV files processed : {len(csv_files)}")
    print(f"Converted to datetime : {converted_count}")
    print(f"Left as text (unparsed) : {fallback_count}")

    stats = {
        "rows": len(df),
        "files": len(csv_files),
        "converted": converted_count,
        "fallback": fallback_count,
        "warning": row_count_warning,
    }

    return final_buffer.getvalue(), stats


# =========================
# STREAMLIT UI
# =========================
st.title("Ticket Automation")

uploaded_files = st.file_uploader(
    "Upload CSV files",
    type=["csv"],
    accept_multiple_files=True,
)

process_clicked = st.button("Process", type="primary")

if process_clicked:
    # Reset any previous result
    st.session_state.pop("excel_bytes", None)
    st.session_state.pop("excel_name", None)
    st.session_state.pop("stats", None)

    if not uploaded_files:
        st.error("No file uploaded. Please upload at least one CSV file.")
    else:
        try:
            with st.spinner("Processing files..."):
                excel_bytes, stats = process_uploaded_files(uploaded_files)

            st.session_state["excel_bytes"] = excel_bytes
            st.session_state["excel_name"] = build_output_filename()
            st.session_state["stats"] = stats

        except MissingShapefileError as e:
            st.error("Missing shapefile")
            st.code(str(e))
        except InvalidCSVError as e:
            st.error(f"Invalid CSV: {e}")
        except ValueError as e:
            st.error(f"Invalid input: {e}")
        except Exception as e:
            st.error(f"An unexpected error occurred while processing: {e}")
            st.exception(e)

# Show result + download button if processing succeeded
if "excel_bytes" in st.session_state:
    stats = st.session_state["stats"]

    st.success("Processing completed successfully!")

    if stats.get("warning"):
        st.warning(stats["warning"])

    st.write(f"Rows processed : {stats['rows']}")
    st.write(f"CSV files processed : {stats['files']}")
    st.write(f"Converted to datetime : {stats['converted']}")
    st.write(f"Left as text (unparsed) : {stats['fallback']}")

    st.download_button(
        label="Download Excel",
        data=st.session_state["excel_bytes"],
        file_name=st.session_state["excel_name"],
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
